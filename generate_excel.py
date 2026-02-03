import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def create_badminton_schedule():
    wb = openpyxl.Workbook()
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']

    ws = wb.create_sheet(title="Vòng Bảng")

    # Group A Teams
    group_a_teams = ["Đội 1", "Đội 4", "Đội 5", "Đội 6", "Đội 9"]
    matches_a = [
        ("Đội 1", "Đội 5"), ("Đội 5", "Đội 9"),
        ("Đội 4", "Đội 6"), ("Đội 1", "Đội 9"),
        ("Đội 5", "Đội 4"), ("Đội 9", "Đội 6"),
        ("Đội 1", "Đội 4"), ("Đội 5", "Đội 6"),
        ("Đội 9", "Đội 4"), ("Đội 1", "Đội 6")
    ]

    # Group B Teams
    group_b_teams = ["Đội 2", "Đội 3", "Đội 7", "Đội 8"]
    matches_b = [
        ("Đội 7", "Đội 3"),
        ("Đội 3", "Đội 8"),
        ("Đội 8", "Đội 2"),
        ("Đội 7", "Đội 8"),
        ("Đội 3", "Đội 2"),
        ("Đội 7", "Đội 2") 
    ]

    # Interleave Matches: 2 A, 1 B pattern
    final_schedule = []
    idx_a = 0
    idx_b = 0
    
    while idx_a < len(matches_a) or idx_b < len(matches_b):
        # Add up to 2 matches from A
        for _ in range(2):
            if idx_a < len(matches_a):
                final_schedule.append(matches_a[idx_a] + ("A",))
                idx_a += 1
        
        # Add 1 match from B
        if idx_b < len(matches_b):
            final_schedule.append(matches_b[idx_b] + ("B",))
            idx_b += 1

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # --- Match Schedule Section ---
    ws.merge_cells('A1:J1')
    ws['A1'] = "Lịch Thi Đấu & Kết Quả Vòng Bảng (1 Set)"
    ws['A1'].font = Font(size=14, bold=True)
    ws['A1'].alignment = center_align

    # Layout:
    # A: Match ID
    # B: Group
    # C: Matchup (Visible)
    # D: P1 Name (Hidden)
    # E: P2 Name (Hidden)
    # F: Set 1 P1 Score (Input)
    # G: Set 1 P2 Score (Input)
    # H: Winner (Formula)
    # I: Loser (Formula)
    # J: Score Display (Formula)

    col_mapping = {
        1: "Trận",
        2: "Bảng",
        3: "Cặp Đấu",
        6: "Điểm Đ1",
        7: "Điểm Đ2",
        8: "Đội Thắng",
        9: "Đội Thua",
        10: "Tỷ Số",
        11: "Hiệu Số"
    }

    header_row_idx = 3
    for col_idx, header_text in col_mapping.items():
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.value = header_text
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    # Hidden Headers
    ws.cell(row=header_row_idx, column=4, value="Ref P1").font = Font(color="CCCCCC")
    ws.cell(row=header_row_idx, column=5, value="Ref P2").font = Font(color="CCCCCC")

    current_row = 4
    match_counter = 1
    matches_start_row = current_row

    for p1, p2, group in final_schedule:
        _add_match_row(ws, current_row, match_counter, group, p1, p2, center_align, thin_border)
        current_row += 1
        match_counter += 1

    last_match_row = current_row - 1
    
    # Hide helper columns
    ws.column_dimensions['D'].hidden = True
    ws.column_dimensions['E'].hidden = True
    
    # Dimensions
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 6
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['H'].width = 15
    ws.column_dimensions['I'].width = 15
    ws.column_dimensions['J'].width = 10
    ws.column_dimensions['K'].width = 10

    # Ranges for Standings
    range_p1 = f"$D${matches_start_row}:$D${last_match_row}"
    range_p2 = f"$E${matches_start_row}:$E${last_match_row}"
    range_s1_p1 = f"$F${matches_start_row}:$F${last_match_row}"
    range_s1_p2 = f"$G${matches_start_row}:$G${last_match_row}"
    range_winner = f"$H${matches_start_row}:$H${last_match_row}"

    # --- Standings Section ---
    standings_start_col = 12 # L

    current_st_row = 3
    current_st_row = create_standing_table(ws, current_st_row, "Bảng Xếp Hạng - Bảng A", group_a_teams, 
                                           standings_start_col, header_font, header_fill, center_align, thin_border, 
                                           range_p1, range_p2, range_s1_p1, range_s1_p2, range_winner)
    
    current_st_row = create_standing_table(ws, current_st_row, "Bảng Xếp Hạng - Bảng B", group_b_teams, 
                                           standings_start_col, header_font, header_fill, center_align, thin_border,
                                           range_p1, range_p2, range_s1_p1, range_s1_p2, range_winner)

    ws.column_dimensions[get_column_letter(standings_start_col+1)].width = 20

    fn = "Badminton_Group_Stage_Final.xlsx"
    wb.save(fn)
    print(f"File created: {fn}")

def _add_match_row(ws, r, match_num, group, p1, p2, align, border):
    ws.cell(row=r, column=1, value=f"Trận {match_num}")
    ws.cell(row=r, column=2, value=group)
    ws.cell(row=r, column=4, value=p1)
    ws.cell(row=r, column=5, value=p2)

    p1_ref_hidden = f"D{r}"
    p2_ref_hidden = f"E{r}"
    ws.cell(row=r, column=3, value=f'={p1_ref_hidden} & " vs " & {p2_ref_hidden}')

    ws.cell(row=r, column=6).alignment = align
    ws.cell(row=r, column=7).alignment = align
    
    s1_p1 = f"F{r}"
    s1_p2 = f"G{r}"

    # Winner
    ws.cell(row=r, column=8, value=f'=IF({s1_p1}>{s1_p2}, {p1_ref_hidden}, IF({s1_p2}>{s1_p1}, {p2_ref_hidden}, ""))')
    
    # Loser
    ws.cell(row=r, column=9, value=f'=IF({s1_p1}>{s1_p2}, {p2_ref_hidden}, IF({s1_p2}>{s1_p1}, {p1_ref_hidden}, ""))')

    # Score Display
    ws.cell(row=r, column=10, value=f'={s1_p1} & "-" & {s1_p2}')
    
    # Point Diff (Hiệu Số)
    # ABS(Score1 - Score2)
    ws.cell(row=r, column=11, value=f'=ABS({s1_p1}-{s1_p2})')

    for c in [1, 2, 3, 6, 7, 8, 9, 10, 11]:
        cell = ws.cell(row=r, column=c)
        cell.border = border
        cell.alignment = align

def create_standing_table(ws, start_row, title, group_players, start_col, h_font, h_fill, align, border, r_p1, r_p2, r_s1_p1, r_s1_p2, r_winner):
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col+6)
    ws.cell(row=start_row, column=start_col, value=title).font = Font(bold=True, size=12)
    ws.cell(row=start_row, column=start_col).alignment = align
    
    st_headers = ["Hạng", "Đội", "Số Trận", "Thắng", "Thua", "Hiệu Số", "Điểm"]
    h_row = start_row + 1
    for i, h in enumerate(st_headers):
        c_idx = start_col + i
        c = ws.cell(row=h_row, column=c_idx)
        c.value = h
        c.font = h_font
        c.fill = h_fill
        c.alignment = align
        c.border = border
        
    data_start_row = h_row + 1
    for i, player in enumerate(group_players):
        r = data_start_row + i
        c = start_col
        
        # Rank: Based on Points (col c+6)
        pts_ref = f"{get_column_letter(c+6)}{r}"
        pts_range = f"{get_column_letter(c+6)}{data_start_row}:{get_column_letter(c+6)}{data_start_row + len(group_players) - 1}"
        ws.cell(row=r, column=c, value=f'=RANK({pts_ref}, {pts_range})')

        # Team
        ws.cell(row=r, column=c+1, value=player)
        player_ref_for_calc = f"{get_column_letter(c+1)}{r}"

        # Matches Played
        ws.cell(row=r, column=c+2, value=f'=COUNTIF({r_p1}, {player_ref_for_calc}) + COUNTIF({r_p2}, {player_ref_for_calc})')

        # Won
        ws.cell(row=r, column=c+3, value=f'=COUNTIF({r_winner}, {player_ref_for_calc})')

        # Lost
        match_played_ref = f"{get_column_letter(c+2)}{r}"
        won_ref = f"{get_column_letter(c+3)}{r}"
        ws.cell(row=r, column=c+4, value=f'={match_played_ref}-{won_ref}')

        # Hiệu Số (Point Diff)
        # Sum(Score when P1) + Sum(Score when P2) - Sum(Conceded when P1) - Sum(Conceded when P2)
        # Points Scored as P1: SUMIF(r_p1, player, r_s1_p1)
        # Points Scored as P2: SUMIF(r_p2, player, r_s1_p2)
        # Points Conceded as P1: SUMIF(r_p1, player, r_s1_p2)
        # Points Conceded as P2: SUMIF(r_p2, player, r_s1_p1)
        
        scored = f'(SUMIF({r_p1}, {player_ref_for_calc}, {r_s1_p1}) + SUMIF({r_p2}, {player_ref_for_calc}, {r_s1_p2}))'
        conceded = f'(SUMIF({r_p1}, {player_ref_for_calc}, {r_s1_p2}) + SUMIF({r_p2}, {player_ref_for_calc}, {r_s1_p1}))'
        
        ws.cell(row=r, column=c+5, value=f'={scored} - {conceded}')

        # Points
        ws.cell(row=r, column=c+6, value=f'={won_ref}*1')

        # Styling
        for col_offset in range(7):
            cell = ws.cell(row=r, column=c+col_offset)
            cell.border = border
            cell.alignment = align

    return data_start_row + len(group_players) + 2

if __name__ == "__main__":
    create_badminton_schedule()
